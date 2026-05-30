"""Bare-bones cashflow tracker: just headers + empty rows."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

HEADERS = ['日期', '每月股息', 'Bitfinex 放贷', '港股打新', '链上机会', '守拙基金分红']
NUM_FMT_USD = '"$"#,##0.##;[Red]"-$"#,##0.##;""'

wb = Workbook()
ws = wb.active
ws.title = '现金流'

# 表头
for i, h in enumerate(HEADERS, start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(name='Helvetica Neue', size=11, bold=True)
    c.alignment = Alignment(horizontal='left', vertical='center')

# 列宽
ws.column_dimensions['A'].width = 14
for i in range(2, len(HEADERS) + 1):
    ws.column_dimensions[chr(64 + i)].width = 16

# 数字格式
for r in range(2, 200):
    ws.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    for c in range(2, len(HEADERS) + 1):
        ws.cell(row=r, column=c).number_format = NUM_FMT_USD

# 冻结表头
ws.freeze_panes = 'A2'

out = '/home/user/cashflow-lab/google-sheets/cashflow-plain.xlsx'
wb.save(out)
print(f'Saved: {out}')
