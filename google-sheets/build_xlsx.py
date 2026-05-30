"""Generate cashflow tracker as a styled .xlsx file."""
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, NamedStyle,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, PieChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout

# ============ Theme ============
BG       = 'F5F5F7'
CARD     = 'FFFFFF'
BORDER   = 'EBEBEB'
TEXT1    = '1A1A1A'
TEXT2    = '666666'
TEXT3    = 'BBBBBB'
ACCENT   = 'FA6400'
POSITIVE = '09B87A'
INP_BG   = 'FFFBF5'
TH_BG    = 'F8F8F8'
TOTAL_BG = 'F0FBF5'

CATEGORIES = ['每月股息', 'Bitfinex 放贷', '港股打新', '链上机会', '守拙基金分红']
RULES = ['不炒股票', '不炒币，只做周期交易', '不买山寨，远离诈骗']
YEAR = datetime.now().year

NUM_FMT_USD = '"$"#,##0;[Red]"-$"#,##0;"—"'
NUM_FMT_PCT = '0.0%;[Red]-0.0%;"—"'

thin = Side(style='thin', color=BORDER)

def hdr_style(cell):
    cell.font = Font(name='Helvetica Neue', size=11, bold=True, color=TEXT2)
    cell.fill = PatternFill('solid', fgColor=TH_BG)
    cell.alignment = Alignment(horizontal='left', vertical='center')
    cell.border = Border(bottom=thin)

def body_font(cell, color=TEXT1, size=12, bold=False, align='right'):
    cell.font = Font(name='Helvetica Neue', size=size, color=color, bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical='center')

def total_style(cell, is_first=False):
    cell.font = Font(name='Helvetica Neue', size=12, bold=True,
                     color=TEXT1 if is_first else POSITIVE)
    cell.fill = PatternFill('solid', fgColor=TOTAL_BG)
    cell.alignment = Alignment(horizontal='left' if is_first else 'right',
                               vertical='center')
    cell.border = Border(top=thin)

wb = Workbook()
default = wb.active
wb.remove(default)

# ============ 设置 ============
ws_settings = wb.create_sheet('设置')
ws_settings['A1'] = '项'
ws_settings['B1'] = '值'
hdr_style(ws_settings['A1']); hdr_style(ws_settings['B1'])

settings_rows = [
    ('年份', YEAR),
    ('货币', 'USD'),
    ('', ''),
    ('纪律红线 1', RULES[0]),
    ('纪律红线 2', RULES[1]),
    ('纪律红线 3', RULES[2]),
]
for i, (k, v) in enumerate(settings_rows, start=2):
    ws_settings.cell(row=i, column=1, value=k)
    ws_settings.cell(row=i, column=2, value=v)
    body_font(ws_settings.cell(row=i, column=1), color=TEXT2, align='left')
    body_font(ws_settings.cell(row=i, column=2), align='left')

ws_settings.column_dimensions['A'].width = 22
ws_settings.column_dimensions['B'].width = 36
ws_settings.row_dimensions[1].height = 26
for r in range(2, 9):
    ws_settings.row_dimensions[r].height = 24
ws_settings.sheet_view.showGridLines = False

# ============ 类别 ============
ws_cat = wb.create_sheet('类别')
ws_cat['A1'] = '类别名'
ws_cat['B1'] = '类型'
hdr_style(ws_cat['A1']); hdr_style(ws_cat['B1'])
for i, name in enumerate(CATEGORIES, start=2):
    ws_cat.cell(row=i, column=1, value=name)
    ws_cat.cell(row=i, column=2, value='收入')
    body_font(ws_cat.cell(row=i, column=1), align='left')
    body_font(ws_cat.cell(row=i, column=2), color=TEXT2, align='left')

dv_type = DataValidation(type='list', formula1='"收入,定投"', allow_blank=True)
dv_type.add(f'B2:B100')
ws_cat.add_data_validation(dv_type)

ws_cat.column_dimensions['A'].width = 22
ws_cat.column_dimensions['B'].width = 14
ws_cat.row_dimensions[1].height = 26
for r in range(2, 20):
    ws_cat.row_dimensions[r].height = 24
ws_cat.sheet_view.showGridLines = False

# ============ 录入 ============
ws_entries = wb.create_sheet('录入')
headers = ['日期', '月份', '类别', '金额(USD)', '备注']
for i, h in enumerate(headers, start=1):
    c = ws_entries.cell(row=1, column=i, value=h)
    hdr_style(c)

# 类别下拉
dv_cat = DataValidation(type='list', formula1=f'=类别!$A$2:$A$100', allow_blank=True)
dv_cat.add('C2:C1000')
ws_entries.add_data_validation(dv_cat)

# 数字格式
for r in range(2, 1001):
    ws_entries.cell(row=r, column=1).number_format = 'yyyy-mm-dd'
    ws_entries.cell(row=r, column=2).number_format = '@'
    ws_entries.cell(row=r, column=4).number_format = NUM_FMT_USD
    ws_entries.cell(row=r, column=4).fill = PatternFill('solid', fgColor=INP_BG)
    for c in range(1, 6):
        body_font(ws_entries.cell(row=r, column=c),
                  align='right' if c == 4 else 'left')
    ws_entries.row_dimensions[r].height = 24

widths = [14, 12, 18, 16, 32]
for i, w in enumerate(widths, start=1):
    ws_entries.column_dimensions[get_column_letter(i)].width = w
ws_entries.row_dimensions[1].height = 28
ws_entries.freeze_panes = 'A2'
ws_entries.sheet_view.showGridLines = False

# ============ 本金 ============
ws_p = wb.create_sheet('本金')
p_headers = ['类别', '本金 (USD)', 'YTD 收入', '年化']
for i, h in enumerate(p_headers, start=1):
    hdr_style(ws_p.cell(row=1, column=i, value=h))

for i, cat in enumerate(CATEGORIES):
    r = i + 2
    ws_p.cell(row=r, column=1, value=cat)
    ws_p.cell(row=r, column=2, value=0)
    ws_p.cell(row=r, column=3, value=(
        f'=IFERROR(SUMIFS(录入!D:D,录入!C:C,A{r},'
        f'录入!B:B,">="&设置!B2&"-01",录入!B:B,"<="&设置!B2&"-12"),0)'
    ))
    ws_p.cell(row=r, column=4, value=f'=IF(B{r}>0,C{r}/B{r},"")')
    body_font(ws_p.cell(row=r, column=1), align='left')
    body_font(ws_p.cell(row=r, column=2))
    body_font(ws_p.cell(row=r, column=3))
    body_font(ws_p.cell(row=r, column=4), color=ACCENT, bold=True)

    ws_p.cell(row=r, column=2).fill = PatternFill('solid', fgColor=INP_BG)
    ws_p.cell(row=r, column=2).number_format = NUM_FMT_USD
    ws_p.cell(row=r, column=3).number_format = NUM_FMT_USD
    ws_p.cell(row=r, column=4).number_format = NUM_FMT_PCT
    ws_p.row_dimensions[r].height = 28

total_r = len(CATEGORIES) + 2
ws_p.cell(row=total_r, column=1, value='合计')
ws_p.cell(row=total_r, column=2, value=f'=SUM(B2:B{total_r-1})')
ws_p.cell(row=total_r, column=3, value=f'=SUM(C2:C{total_r-1})')
ws_p.cell(row=total_r, column=4, value=f'=IF(B{total_r}>0,C{total_r}/B{total_r},"")')
for c in range(1, 5):
    total_style(ws_p.cell(row=total_r, column=c), is_first=(c == 1))
ws_p.cell(row=total_r, column=2).number_format = NUM_FMT_USD
ws_p.cell(row=total_r, column=3).number_format = NUM_FMT_USD
ws_p.cell(row=total_r, column=4).number_format = NUM_FMT_PCT
ws_p.row_dimensions[total_r].height = 32

for col, w in zip('ABCD', [22, 16, 16, 12]):
    ws_p.column_dimensions[col].width = w
ws_p.row_dimensions[1].height = 28
ws_p.freeze_panes = 'A2'
ws_p.sheet_view.showGridLines = False

# ============ 月度汇总 ============
ws_sum = wb.create_sheet('月度汇总')
sum_headers = ['月份'] + CATEGORIES + ['本月合计', 'YTD']
for i, h in enumerate(sum_headers, start=1):
    hdr_style(ws_sum.cell(row=1, column=i, value=h))

n_cats = len(CATEGORIES)
total_col = n_cats + 2
ytd_col = n_cats + 3
total_col_letter = get_column_letter(total_col)

for m in range(12):
    r = m + 2
    # Month label: =设置!B2 & "-" & TEXT(m+1, "00")
    ws_sum.cell(row=r, column=1, value=f'=设置!B2&"-"&TEXT({m+1},"00")')
    body_font(ws_sum.cell(row=r, column=1), color=TEXT2, align='left')

    for ci in range(n_cats):
        col = ci + 2
        col_letter = get_column_letter(col)
        ws_sum.cell(row=r, column=col, value=(
            f'=IFERROR(SUMIFS(录入!D:D,录入!C:C,{col_letter}$1,'
            f'录入!B:B,$A{r}),0)'
        ))
        ws_sum.cell(row=r, column=col).number_format = NUM_FMT_USD
        body_font(ws_sum.cell(row=r, column=col))

    # 本月合计
    last_cat_letter = get_column_letter(n_cats + 1)
    ws_sum.cell(row=r, column=total_col, value=f'=SUM(B{r}:{last_cat_letter}{r})')
    ws_sum.cell(row=r, column=total_col).number_format = NUM_FMT_USD
    body_font(ws_sum.cell(row=r, column=total_col), color=POSITIVE, bold=True)

    # YTD
    ws_sum.cell(row=r, column=ytd_col,
                value=f'=SUM(${total_col_letter}$2:{total_col_letter}{r})')
    ws_sum.cell(row=r, column=ytd_col).number_format = NUM_FMT_USD
    body_font(ws_sum.cell(row=r, column=ytd_col), color=POSITIVE, bold=True)

    ws_sum.row_dimensions[r].height = 28

# 年度合计行
tr = 14
ws_sum.cell(row=tr, column=1, value='年度合计')
for c in range(2, ytd_col):
    col_l = get_column_letter(c)
    ws_sum.cell(row=tr, column=c, value=f'=SUM({col_l}2:{col_l}13)')
ws_sum.cell(row=tr, column=ytd_col,
            value=f'={total_col_letter}{tr}')
for c in range(1, ytd_col + 1):
    total_style(ws_sum.cell(row=tr, column=c), is_first=(c == 1))
    if c >= 2:
        ws_sum.cell(row=tr, column=c).number_format = NUM_FMT_USD
ws_sum.row_dimensions[tr].height = 32

ws_sum.column_dimensions['A'].width = 12
for c in range(2, ytd_col + 1):
    ws_sum.column_dimensions[get_column_letter(c)].width = 16
ws_sum.row_dimensions[1].height = 28
ws_sum.freeze_panes = 'B2'
ws_sum.sheet_view.showGridLines = False

# ============ 仪表盘 ============
ws_dash = wb.create_sheet('仪表盘', 0)
ws_dash.sheet_view.showGridLines = False

# 标题
ws_dash['B2'] = 'CASHFLOW TRACKER'
ws_dash['B2'].font = Font(name='Helvetica Neue', size=10, color=TEXT3,
                          bold=False)
ws_dash['B2'].alignment = Alignment(horizontal='left', vertical='center')

ws_dash['B3'] = '现金流 · 跟踪表'
ws_dash['B3'].font = Font(name='Helvetica Neue', size=22, bold=True,
                          color=TEXT1)
ws_dash['B3'].alignment = Alignment(horizontal='left', vertical='center')

ws_dash['B4'] = f'="始于 "&设置!B2&" 年 · 月月不落"'
ws_dash['B4'].font = Font(name='Helvetica Neue', size=10, color=TEXT3)
ws_dash['B4'].alignment = Alignment(horizontal='left', vertical='center')

# 右上 KPI
ws_dash['G2'] = '年度现金流'
ws_dash['G2'].font = Font(name='Helvetica Neue', size=10, color=TEXT2)
ws_dash['G2'].alignment = Alignment(horizontal='right', vertical='center')

ws_dash['G3'] = f'=月度汇总!{get_column_letter(ytd_col)}{tr}'
ws_dash['G3'].font = Font(name='Helvetica Neue', size=20, bold=True,
                          color=POSITIVE)
ws_dash['G3'].alignment = Alignment(horizontal='right', vertical='center')
ws_dash['G3'].number_format = NUM_FMT_USD
ws_dash['G3'].fill = PatternFill('solid', fgColor=TOTAL_BG)

# KPI 卡片
kpi_row = 7
kpi_labels = ['月均', '本月', '总本金', '年化']
for i, label in enumerate(kpi_labels):
    col = i + 2
    c_label = ws_dash.cell(row=kpi_row, column=col, value=label)
    c_label.font = Font(name='Helvetica Neue', size=10, color=TEXT2)
    c_label.alignment = Alignment(horizontal='left', vertical='center')
    c_label.fill = PatternFill('solid', fgColor=CARD)
    c_label.border = Border(left=thin, right=thin, top=thin)

month_count_part = f'IF(MONTH(TODAY())=0,1,MONTH(TODAY()))'
ws_dash.cell(row=kpi_row + 1, column=2,
             value=f'=IFERROR(月度汇总!{get_column_letter(ytd_col)}{tr}/{month_count_part},0)')
ws_dash.cell(row=kpi_row + 1, column=3,
             value=f'=IFERROR(INDEX(月度汇总!{total_col_letter}2:{total_col_letter}13,MONTH(TODAY())),0)')
ws_dash.cell(row=kpi_row + 1, column=4,
             value=f'=本金!B{total_r}')
ws_dash.cell(row=kpi_row + 1, column=5,
             value=f'=IF(D{kpi_row+1}>0,月度汇总!{get_column_letter(ytd_col)}{tr}/D{kpi_row+1},"")')

for i in range(4):
    col = i + 2
    c = ws_dash.cell(row=kpi_row + 1, column=col)
    c.font = Font(name='Helvetica Neue', size=18, bold=True, color=TEXT1)
    c.alignment = Alignment(horizontal='left', vertical='center')
    c.fill = PatternFill('solid', fgColor=CARD)
    c.border = Border(left=thin, right=thin, bottom=thin)
    c.number_format = NUM_FMT_PCT if i == 3 else NUM_FMT_USD

# 列宽 / 行高
ws_dash.column_dimensions['A'].width = 3
for c in range(2, 8):
    ws_dash.column_dimensions[get_column_letter(c)].width = 18
ws_dash.column_dimensions['H'].width = 3

ws_dash.row_dimensions[2].height = 18
ws_dash.row_dimensions[3].height = 32
ws_dash.row_dimensions[4].height = 18
ws_dash.row_dimensions[kpi_row].height = 20
ws_dash.row_dimensions[kpi_row + 1].height = 36

# 区块标题
ws_dash['B11'] = '月度趋势'
ws_dash['B11'].font = Font(name='Helvetica Neue', size=11, bold=True,
                           color=TEXT1)
ws_dash['B27'] = '类别占比'
ws_dash['B27'].font = Font(name='Helvetica Neue', size=11, bold=True,
                           color=TEXT1)
ws_dash['E27'] = '月度构成'
ws_dash['E27'].font = Font(name='Helvetica Neue', size=11, bold=True,
                           color=TEXT1)

# ============ 图表 ============
# 1. 月度趋势线图
line = LineChart()
line.title = None
line.legend = None
line.height = 8
line.width = 24
data_ref = Reference(ws_sum, min_col=total_col, min_row=1,
                     max_col=total_col, max_row=13)
cat_ref = Reference(ws_sum, min_col=1, min_row=2, max_row=13)
line.add_data(data_ref, titles_from_data=True)
line.set_categories(cat_ref)
if line.series:
    s = line.series[0]
    s.graphicalProperties.line.solidFill = ACCENT
    s.graphicalProperties.line.width = 20000
    if s.marker is None:
        from openpyxl.chart.marker import Marker
        s.marker = Marker(symbol='circle', size=5)
        s.marker.graphicalProperties.solidFill = ACCENT
ws_dash.add_chart(line, 'B12')

# 2. 类别占比饼图（用年度合计行）
pie = PieChart()
pie.title = None
pie.height = 7
pie.width = 10
pie_data = Reference(ws_sum, min_col=2, min_row=tr, max_col=n_cats + 1,
                     max_row=tr)
pie_labels = Reference(ws_sum, min_col=2, min_row=1, max_col=n_cats + 1,
                       max_row=1)
pie.add_data(pie_data, titles_from_data=False)
pie.set_categories(pie_labels)
pie.dataLabels = DataLabelList(showPercent=True)
ws_dash.add_chart(pie, 'B28')

# 3. 月度堆叠柱图
bar = BarChart()
bar.type = 'col'
bar.grouping = 'stacked'
bar.overlap = 100
bar.title = None
bar.height = 7
bar.width = 13
bar_data = Reference(ws_sum, min_col=2, min_row=1,
                     max_col=n_cats + 1, max_row=13)
bar_cat = Reference(ws_sum, min_col=1, min_row=2, max_row=13)
bar.add_data(bar_data, titles_from_data=True)
bar.set_categories(bar_cat)
ws_dash.add_chart(bar, 'E28')

# ============ 排序 sheets ============
order = ['仪表盘', '录入', '类别', '本金', '月度汇总', '设置']
wb._sheets = [wb[name] for name in order]

# ============ 保存 ============
out = '/home/user/cashflow-lab/google-sheets/cashflow-tracker.xlsx'
wb.save(out)
print(f'Saved: {out}')
